"""Drawing primitives shared by the P3 deck. Same visual language as the P2 deck:
cream background, orange/green accents, Georgia headings, native shapes and charts only.
Generated files go to P3/ at the repository root, which Git ignores."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "output" / "friendship_gpa"
P3 = ROOT / "P3"
P3.mkdir(exist_ok=True)
W, H = 13.333, 7.5
BG, WHITE, ORANGE, DARK, GREEN, GREY, LINE, SAND, NODE = (
    "FAF8F5", "FFFFFF", "C4602F", "262522", "4F7A5F", "6F6C67", "E3DDD3", "F5F1EA", "BDB8B0")
HEAD, BODY = "Georgia", "Arial"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
rgb = RGBColor.from_string


def new_slide():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(BG)
    return s


def box(s, x, y, w, h, fill=WHITE, line=LINE, kind=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.09, lw=1):
    sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE: sh.adjustments[0] = adj
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb = rgb(line); sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    return sh


def text(s, x, y, w, h, value, size=14, color=DARK, bold=False, font=BODY, align=PP_ALIGN.LEFT,
         italic=False, anchor=MSO_ANCHOR.TOP, spacing=None):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.04); tf.margin_top = tf.margin_bottom = Inches(0.02)
    for i, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing: p.line_spacing = spacing
        r = p.add_run(); r.text = line
        f = r.font; f.name = font; f.size = Pt(size); f.bold = bold; f.italic = italic; f.color.rgb = rgb(color)
    return tb


def line(s, x1, y1, x2, y2, color=NODE, w=1.25, dash=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color); c.line.width = Pt(w)
    if dash: c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return c


def node(s, x, y, label="", fill=GREY, r=0.2, size=10, color=WHITE):
    """Circle centred on (x, y) with an optional label inside."""
    box(s, x - r, y - r, 2 * r, 2 * r, fill, WHITE, MSO_SHAPE.OVAL, lw=1.5)
    if label:
        text(s, x - r - 0.3, y - 0.15, 2 * r + 0.6, 0.3, label, size, color, True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def network(s, nodes, edges, r=0.2, size=10):
    """nodes: {name: (x, y, label, fill)}; edges: (a, b[, colour, dashed])."""
    for e in edges:
        a, b = nodes[e[0]], nodes[e[1]]
        col = e[2] if len(e) > 2 else NODE
        line(s, a[0], a[1], b[0], b[1], col, 1.75, dash=len(e) > 3 and e[3])
    for x, y, lab, fill in nodes.values():
        node(s, x, y, lab, fill, r, size)


def header(s, kicker, title, two_line=False):
    text(s, 0.79, 0.42, 9, 0.3, kicker.upper(), 10.5, ORANGE, True)
    text(s, 0.79, 0.72, 11.8, 0.7 if not two_line else 1.2, title, 28, DARK, True, HEAD)
    box(s, 0.8, 1.5 if not two_line else 1.95, 0.8, 0.035, ORANGE, None, MSO_SHAPE.RECTANGLE)


def pagenum(s, n):
    text(s, 12.2, 7.03, 0.7, 0.3, str(n), 9, GREY, align=PP_ALIGN.RIGHT)


def callout(s, y, label, msg, h=0.85, x=0.8, w=11.75, size=14):
    box(s, x, y, w, h, SAND, ORANGE, adj=0.12)
    if label:
        text(s, x + 0.3, y, 1.9, h, label, 11.5, ORANGE, True, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 2.3, y, w - 2.6, h, msg, size, DARK, True, anchor=MSO_ANCHOR.MIDDLE)
    else:
        text(s, x + 0.3, y, w - 0.6, h, msg, size, DARK, True, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)


def stat(s, x, y, w, h, label, value, sub, color=DARK, vsize=26):
    box(s, x, y, w, h, WHITE, LINE, adj=0.12)
    text(s, x + 0.2, y + 0.12, w - 0.4, 0.25, label.upper(), 9.5, GREY, True)
    text(s, x + 0.2, y + 0.38, w - 0.4, 0.55, value, vsize, color, True, HEAD)
    text(s, x + 0.2, y + h - 0.32, w - 0.4, 0.25, sub, 9, GREY)


def chip(s, x, y, w, h, top, bottom, col=ORANGE, fill=WHITE):
    box(s, x, y, w, h, fill, LINE, adj=0.14)
    box(s, x, y + 0.14, 0.09, h - 0.28, col, None, MSO_SHAPE.RECTANGLE)
    text(s, x + 0.25, y + 0.08, w - 0.35, 0.3, top, 10, GREY, True)
    text(s, x + 0.25, y + 0.38, w - 0.35, h - 0.45, bottom, 15, DARK, True, HEAD)


def forest(s, x, y, w, rows, lo, hi, step, row_h=0.85, label_w=2.6, note_w=2.0, cap=True, fs=13):
    """rows: (label, est, ci_lo, ci_hi, note, colour)."""
    px, pw = x + label_w, w - label_w - note_w
    sx = lambda v: px + (v - lo) / (hi - lo) * pw
    h = row_h * len(rows)
    for k in range(int(round((hi - lo) / step)) + 1):
        v = lo + k * step
        box(s, sx(v) - 0.004, y, 0.008, h, "E9E4DB", None, MSO_SHAPE.RECTANGLE)
        lab = "0" if abs(v) < 1e-9 else f"{v:+.2f}"
        text(s, sx(v) - 0.4, y + h + 0.04, 0.8, 0.25, lab, 9.5, GREY, align=PP_ALIGN.CENTER)
    box(s, sx(0) - 0.011, y - 0.06, 0.022, h + 0.12, DARK, None, MSO_SHAPE.RECTANGLE)
    for i, (label, est, a, b, note, col) in enumerate(rows):
        cy = y + i * row_h + row_h / 2
        text(s, x, cy - 0.3, label_w - 0.1, 0.6, label, fs, DARK, True, anchor=MSO_ANCHOR.MIDDLE)
        box(s, sx(a), cy - 0.02, sx(b) - sx(a), 0.04, col, None, MSO_SHAPE.RECTANGLE)
        for v in (a, b): box(s, sx(v) - 0.01, cy - 0.1, 0.02, 0.2, col, None, MSO_SHAPE.RECTANGLE)
        box(s, sx(est) - 0.1, cy - 0.1, 0.2, 0.2, col, WHITE, MSO_SHAPE.OVAL)
        text(s, px + pw + 0.1, cy - 0.3, note_w - 0.1, 0.6, note, fs - 1, GREY, anchor=MSO_ANCHOR.MIDDLE)
    if cap:
        text(s, px, y + h + 0.3, pw, 0.25, "Coefficient, 95% confidence interval", 9.5, GREY,
             italic=True, align=PP_ALIGN.CENTER)


def bars(s, x, y, w, items, vmax, row_h=0.55, label_w=2.3, fs=12, val_fmt="{:+.3f}"):
    """items: (label, value, colour[, display-text])."""
    bw = w - label_w - 1.1
    for i, it in enumerate(items):
        lab, v, col = it[0], it[1], it[2]
        yy = y + i * row_h
        text(s, x, yy, label_w - 0.1, row_h - 0.1, lab, fs, DARK, True, anchor=MSO_ANCHOR.MIDDLE)
        bw_px = max(0.04, bw * abs(v) / vmax)
        box(s, x + label_w, yy + 0.08, bw_px, row_h - 0.26, col, None, MSO_SHAPE.RECTANGLE)
        disp = it[3] if len(it) > 3 else val_fmt.format(v)
        text(s, x + label_w + bw_px + 0.08, yy, 1.1, row_h - 0.1, disp, fs, DARK, True, anchor=MSO_ANCHOR.MIDDLE)


def line_chart(s, x, y, w, h, series, xvals, ylab="", title="", xlab=""):
    """series: list of (name, [y...], colour). Simple polyline chart."""
    pad_l, pad_b, pad_t = 0.42, 0.3, 0.25
    px0, px1 = x + pad_l, x + w
    py0, py1 = y + h - pad_b, y + pad_t
    raw_min = min(min(vals) for _, vals, _ in series)
    ymax = max(max(vals) for _, vals, _ in series)
    pad = (ymax - raw_min) * 0.12 or 0.05
    ymin = max(0.0, raw_min - pad) if raw_min >= 0 else raw_min - pad
    ymax = ymax + pad
    if title: text(s, x, y - 0.02, w, 0.25, title, 10.5, ORANGE, True)
    for gy in (0, 0.25, 0.5, 0.75, 1.0):
        yy = py0 - gy * (py0 - py1)
        box(s, px0, yy, px1 - px0, 0.006, "E9E4DB", None, MSO_SHAPE.RECTANGLE)
        val = ymin + gy * (ymax - ymin)
        lab = text(s, x - 0.08, yy - 0.09, pad_l - 0.02, 0.18, f"{val:.0f}", 7.5, GREY, align=PP_ALIGN.RIGHT)
        lab.text_frame.word_wrap = False
    n = len(xvals)
    sx = lambda i: px0 + (px1 - px0) * i / max(1, n - 1)
    sy = lambda v: py0 - (v - ymin) / (ymax - ymin) * (py0 - py1)
    for i, xv in enumerate(xvals):
        text(s, sx(i) - 0.3, py0 + 0.03, 0.6, 0.2, str(xv), 8.5, GREY, align=PP_ALIGN.CENTER)
    for name, vals, col in series:
        pts = [(sx(i), sy(v)) for i, v in enumerate(vals)]
        for (x1, y1), (x2, y2) in zip(pts, pts[1:]):
            line(s, x1, y1, x2, y2, col, 2.25)
        for px_, py_ in pts:
            box(s, px_ - 0.045, py_ - 0.045, 0.09, 0.09, col, WHITE, MSO_SHAPE.OVAL)
    if xlab: text(s, x, py0 + 0.22, w, 0.2, xlab, 8.5, GREY, italic=True, align=PP_ALIGN.CENTER)


def legend(s, x, y, items, fs=10):
    """items: (label, colour)."""
    cx = x
    for lab, col in items:
        box(s, cx, y + 0.03, 0.16, 0.16, col, None, MSO_SHAPE.OVAL)
        tb = text(s, cx + 0.22, y, 1.8, 0.25, lab, fs, DARK, True)
        cx += 0.22 + 0.11 * len(lab) + 0.35
