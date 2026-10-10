"""Optional preview: renders each slide of a .pptx to PNG with macOS Quick Look (qlmanage)
and tiles them four to an image. Usage: python3 render_preview.py <deck.pptx> <out_dir>"""
import sys, subprocess, os
from pptx import Presentation
from PIL import Image
src, out = sys.argv[1], sys.argv[2]; os.makedirs(out, exist_ok=True)
for f in os.listdir(out):  # clear only files this script writes
    if (f.startswith("one") and f[3:5].isdigit()) or (f.startswith("m") and f[1:-4].isdigit() and f.endswith(".png")):
        os.remove(f"{out}/{f}")
n=len(Presentation(src).slides)
for i in range(n):
    p=Presentation(src); lst=p.slides._sldIdLst
    for j,e in enumerate(list(lst)):
        if j!=i: lst.remove(e)
    f=f"{out}/one{i+1:02d}.pptx"; p.save(f)
    subprocess.run(["qlmanage","-t","-s","1400","-o",out,f],capture_output=True)
ims=[Image.open(f"{out}/one{i:02d}.pptx.png").convert("RGB") for i in range(1,n+1)]
w,h=ims[0].size
for k in range(0,n,4):
    grp=ims[k:k+4]; m=Image.new("RGB",(w*2,h*2),"white")
    for i,im in enumerate(grp): m.paste(im,((i%2)*w,(i//2)*h))
    m.save(f"{out}/m{k//4+1}.png")
print(n)
