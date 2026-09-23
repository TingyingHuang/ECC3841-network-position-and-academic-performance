"""Build the editable Presentation 2 speaker script.

Blue runs are the material added or substantially revised after the
identification audit. The Markdown source lives beside this builder.
"""
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "Presentation2_speaker_script_revised.docx"
BLUE = RGBColor(0x1F, 0x5E, 0xAA)

SLIDES = [
    ("Slide 1  Network Position and Academic Performance", [
        ("Good afternoon. Our project asks whether a student's position in a directed friendship network predicts academic performance beyond their direct contacts. ", False),
        ("This presentation is a progress report: we found an initially interesting association, audited the model behind it, and changed the primary analysis when the data could not support the original interpretation.", True),
    ]),
    ("Slide 2  How the analysis changed", [
        ("Our question has stayed the same. We ask whether outgoing network reach predicts later grades after accounting for earlier grades and direct contacts. We first replicated known patterns in the data, then estimated a pooled association between Katz centrality and GPA. ", False),
        ("The important next step was an audit of time order and data availability. That audit changed which result we treat as primary.", True),
    ]),
    ("Slide 3  The replication holds", [
        ("First, we replicated the selection baseline in the school panel. Own past GPA strongly predicts next GPA, with a coefficient of 0.955. Average friend's past GPA is not significant once own past GPA is included, with p equal to 0.557. We also reproduce evidence that newly formed university ties connect students who are more alike than ties that disappear. These checks show that the data and basic construction behave as expected. They do not, by themselves, identify a network-position effect.", False),
    ]),
    ("Slide 4  The first pooled result looked significant", [
        ("Our initial pooled analysis showed a positive Katz-GPA association of 0.089 when Katz centrality entered alone. ", False),
        ("That was a descriptive association across all groups, rather than evidence that network position caused later GPA. After we controlled for direct in-degree and out-degree, the coefficient fell to 0.012 and was not significant, with p equal to 0.760. This suggested that direct links rather than wider network reach explained the initial pattern.", True),
    ]),
    ("Slide 5  The model audit", [
        ("We then found the key identification limitation. University GPA is recorded as a static value, not repeatedly across network snapshots. A university cross-section can show that GPA and position are associated, but it cannot control for baseline achievement or establish that the network came before a later GPA outcome. Therefore, the university comparison cannot answer our central next-period question.", True),
    ]),
    ("Slide 6  The corrected primary test", [
        ("Only the school cohort records GPA repeatedly. We therefore use it for the primary longitudinal model. The outcome is later GPA, and we control for prior GPA, outgoing Katz centrality, direct links, average contact GPA and snapshot effects. The Katz estimate is positive but small, 0.009, with p equal to 0.139 and a 95 percent confidence interval from minus 0.003 to 0.022. We therefore do not find statistically significant evidence that outgoing network reach independently predicts later grades.", True),
    ]),
    ("Slide 7  Robustness across phi", [
        ("Katz centrality depends on the attenuation parameter phi, which determines how much indirect links count. ", False),
        ("The pooled cross-sectional estimates vary with phi. In the corrected school panel, however, the estimate remains statistically non-significant across the tested phi range. This means our conclusion does not depend on choosing one convenient Katz parameter.", True),
    ]),
    ("Slide 8  What we can conclude at this stage", [
        ("The initial association did not survive the model audit. In the data that support a temporal test, we cannot conclude that wider outgoing reach improves later GPA. The university result remains descriptive, and it should not be presented as causal evidence or as a basis for an intervention.", True),
    ]),
    ("Slide 9  Questions for Presentation 3", [
        ("Presentation 3 will move from empirical diagnosis to network economics. We will model why students form links, when information or peer effects could travel through those links, and what incentives shape access to useful contacts. We will then connect that mechanism to a cautious recommendation for a non-expert manager, consistent with the evidence and its limits.", True),
    ]),
    ("Slide 10  Close", [
        ("Thank you. ", False),
        ("Our contribution so far is not a claim that centrality raises grades. It is a corrected and more credible assessment of what this dataset can and cannot establish. We welcome your questions.", True),
    ]),
]


def set_font(run, size=11, bold=False, color=RGBColor(0, 0, 0)):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
styles["Normal"].font.size = Pt(11)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Title"].font.name = "Aptos Display"
styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
styles["Title"].font.size = Pt(22)
styles["Heading 1"].font.name = "Aptos Display"
styles["Heading 1"].font.color.rgb = RGBColor(0, 0, 0)
styles["Heading 1"].font.size = Pt(14)

doc.add_paragraph("ECC3841 Presentation 2 Speaker Script", style="Title")
intro = doc.add_paragraph()
intro.paragraph_format.space_after = Pt(16)
r = intro.add_run("Blue text marks material added or substantially revised after the model audit.")
set_font(r, size=10.5, color=BLUE)

for heading, parts in SLIDES:
    doc.add_paragraph(heading, style="Heading 1")
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(11)
    for body, revised in parts:
        r = p.add_run(body)
        set_font(r, color=BLUE if revised else RGBColor(0, 0, 0))

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run("ECC3841 Network Economics  |  Presentation 2")
set_font(r, size=9, color=RGBColor(90, 90, 90))

doc.save(OUT)
print(f"Saved {OUT}")
