"""
Presentation 1, visual-heavy version. Each content slide is one figure
plus an eyebrow, a short title, and a one-line takeaway.

    python3 build_assets_visual.py     # first, generates the figures
    python3 build_visual_deck.py       # then, assembles the deck
"""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parent
A = ROOT / "assets"

BLUE = RGBColor(0x2A, 0x78, 0xD6)
ORANGE = RGBColor(0xEB, 0x68, 0x34)
DARK = RGBColor(0x19, 0x18, 0x0F)
GRAY = RGBColor(0x55, 0x50, 0x3F)
BGCOL = RGBColor(0xFA, 0xF9, 0xF6)
CARD = RGBColor(0xF1, 0xEF, 0xE8)
BORDER = RGBColor(0xDD, 0xD8, 0xCD)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

HEAD = "Georgia"
BODY = "Calibri"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
W = 13.333


def slide():
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    r.fill.solid(); r.fill.fore_color.rgb = BGCOL
    r.line.fill.background(); r.shadow.inherit = False
    s.shapes._spTree.remove(r._element)
    s.shapes._spTree.insert(2, r._element)
    return s


def text(s, left, top, width, height, body, size=18, bold=False, color=DARK,
         font=BODY, align=PP_ALIGN.LEFT, italic=False, spacing=1.12,
         anchor=MSO_ANCHOR.TOP):
    box = s.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, ln in enumerate(body.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ln
        p.alignment = align
        p.line_spacing = spacing
        for run in p.runs:
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.italic = italic
            run.font.color.rgb = color
            run.font.name = font
    return box


def eyebrow(s, t, color=ORANGE):
    text(s, 0.7, 0.42, 11.5, 0.4, t, size=13, bold=True, color=color)


def title(s, t, size=26):
    text(s, 0.7, 0.8, 11.9, 1.0, t, size=size, bold=True, color=DARK, font=HEAD, spacing=1.05)


def takeaway(s, t, color=DARK, bold=True):
    text(s, 0.7, 6.75, 11.9, 0.6, t, size=15, bold=bold, color=color)


def figure(s, name, top=1.62, height=4.95):
    pic = s.shapes.add_picture(str(A / name), 0, Inches(top), height=Inches(height))
    pic.left = int((prs.slide_width - pic.width) / 2)
    return pic


def pagenum(s, n):
    text(s, 12.5, 7.08, 0.7, 0.32, str(n), size=11, color=GRAY, align=PP_ALIGN.RIGHT)


# ===================================================== 1 title
s = slide()
try:
    p = s.shapes.add_picture(str(A / "decor_network.png"), Inches(8.9), Inches(0.4),
                             height=Inches(4.0))
except Exception:
    pass
text(s, 0.9, 2.05, 11.5, 0.5, "ECC3841 NETWORK ECONOMICS  ·  PRESENTATION 1",
     size=13, bold=True, color=ORANGE)
text(s, 0.9, 2.6, 11.5, 1.9, "Network Position and\nAcademic Performance",
     size=42, bold=True, color=DARK, font=HEAD, spacing=1.05)
text(s, 0.9, 4.7, 11.2, 1.0,
     "Does where you sit in a friendship network predict your grades,\n"
     "beyond who your direct friends are?",
     size=18, color=GRAY, italic=True, spacing=1.3)
text(s, 0.9, 6.5, 9, 0.5, "Amy Bing   ·   Tingying Huang", size=13, color=GRAY)

# ===================================================== 2 puzzle
s = slide()
eyebrow(s, "THE PUZZLE")
title(s, "Same friends. Same grades?")
figure(s, "fig_puzzle.png", top=1.55, height=4.9)
takeaway(s, "Research question:  does a student's position in the wider network affect their grades?",
         color=ORANGE)
pagenum(s, 2)

# ===================================================== 3 data
s = slide()
eyebrow(s, "THE DATA")
title(s, "A real friendship network, watched over time")
figure(s, "fig_data.png", top=1.55, height=4.9)
takeaway(s, "Directed “like” ties  ·  ~6,200 Russian students  ·  the same dataset as Smirnov & Thurner (2017).",
         bold=False, color=GRAY)
pagenum(s, 3)

# ===================================================== 4 prior work
s = slide()
eyebrow(s, "PRIOR WORK")
title(s, "Smirnov & Thurner (2017): the effect is mostly selection")
figure(s, "fig_st_formula.png", top=1.65, height=1.35)
figure(s, "fig_priorwork.png", top=3.2, height=3.5)
takeaway(s,
         "One number per student. It never asks where that student sits in the wider network.",
         color=ORANGE)
pagenum(s, 4)

# ===================================================== 5 the idea
s = slide()
eyebrow(s, "THE IDEA WE TEST")
title(s, "Position: how far you reach, not just who you know")
figure(s, "fig_idea.png", top=1.55, height=4.9)
takeaway(s, "We test whether Katz-Bonacich centrality predicts GPA on top of popularity and friend quality.")
pagenum(s, 5)

# ===================================================== 6 plan
s = slide()
eyebrow(s, "OUR PLAN")
title(s, "Three steps")
steps = [
    ("REPLICATE", "Re-run the selection-vs-influence checks on the same data."),
    ("ISOLATE", "Compute Katz-Bonacich across a range of φ; does it beat raw popularity?"),
    ("COMPARE", "High school vs. university — does the effect hold in both?"),
]
x = 0.7
bw, gap = 3.9, 0.42
for tag, desc in steps:
    c = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.5), Inches(bw), Inches(3.0))
    c.adjustments[0] = 0.06
    c.fill.solid(); c.fill.fore_color.rgb = WHITE
    c.line.color.rgb = BORDER; c.line.width = Pt(1.25); c.shadow.inherit = False
    text(s, x + 0.3, 2.85, bw - 0.6, 0.5, tag, size=16, bold=True, color=ORANGE)
    text(s, x + 0.3, 3.55, bw - 0.6, 1.8, desc, size=14, color=GRAY, spacing=1.3)
    x += bw + gap
takeaway(s, "One headline test, at φ = 0.85 × (1/λmax), is fixed in advance.",
         bold=False, color=GRAY)
pagenum(s, 6)

# ===================================================== 7 how
s = slide()
eyebrow(s, "HOW WE'LL GO ABOUT IT")
title(s, "Strip away the alternatives → the structural effect")
figure(s, "fig_filter.png", top=1.55, height=4.95)
takeaway(s, "Does position still predict grades, net of past GPA, friend quality, and popularity?",
         color=ORANGE)
pagenum(s, 7)

# ===================================================== 8 identification
s = slide()
eyebrow(s, "IDENTIFICATION")
title(s, "Which way does the arrow run?")
figure(s, "fig_identification.png", top=1.75, height=4.4)
takeaway(s, "Our claim: position predicts later grades — not that it causes them.", bold=False, color=GRAY)
pagenum(s, 8)

# ===================================================== 9 predictions
s = slide()
eyebrow(s, "PREDICTIONS")
title(s, "What we expect — and what would prove us wrong")
figure(s, "fig_predictions.png", top=1.85, height=4.3)
takeaway(s,
         "Survives the controls → the simple measure misses something.   "
         "Drops to zero → the simple story wins.", bold=False, color=GRAY)
pagenum(s, 9)

# ===================================================== 10 thanks
s = slide()
text(s, 0.9, 2.6, 11, 1.2, "Thank you", size=42, bold=True, color=DARK, font=HEAD)
text(s, 0.95, 4.0, 11, 0.6,
     "The dataset and pipeline are built — results in Presentation 2.", size=16, color=GRAY)
text(s, 0.95, 5.0, 11, 1.0,
     "Data — Smirnov & Thurner (2017), Harvard Dataverse  doi:10.7910/DVN/SZA9YW\n"
     "Model — Ballester, Calvó-Armengol & Zenou (2006); Calvó-Armengol, Patacchini & Zenou (2009)",
     size=13, color=GRAY, spacing=1.4)
text(s, 0.95, 6.4, 6, 0.5, "Questions?", size=16, bold=True, color=ORANGE)
pagenum(s, 10)

out = ROOT / "Presentation1_visual.pptx"
prs.save(out)
print(f"Saved -> {out}  ({len(prs.slides._sldIdLst)} slides)")
